from __future__ import annotations

import tempfile
from pathlib import Path

from .common import create_watermark_layer, find_font
from .constants import EMU_PER_INCH
from .options import WatermarkOptions


def watermark_pptx(input_path: Path, output_path: Path, options: WatermarkOptions) -> Path:
    try:
        from pptx import Presentation
    except ImportError as e:
        raise ImportError("处理 PPTX 需要安装 python-pptx：pip install python-pptx") from e

    prs = Presentation(str(input_path))
    slide_width = prs.slide_width
    slide_height = prs.slide_height

    width_px = max(1, int((slide_width / EMU_PER_INCH) * options.dpi))
    height_px = max(1, int((slide_height / EMU_PER_INCH) * options.dpi))

    font = find_font(options.font_path)
    layer = create_watermark_layer(width_px, height_px, options, font)

    with tempfile.TemporaryDirectory() as td:
        wm_path = Path(td) / "watermark.png"
        layer.save(wm_path, "PNG", dpi=(options.dpi, options.dpi))

        for slide in prs.slides:
            pic = slide.shapes.add_picture(str(wm_path), 0, 0, width=slide_width, height=slide_height)
            sp_tree = slide.shapes._spTree
            pic_el = pic._element
            sp_tree.remove(pic_el)
            sp_tree.insert(2, pic_el)

    prs.save(str(output_path))
    return output_path
