from __future__ import annotations

import re
import tempfile
import zipfile
from pathlib import Path

from .common import (
    add_png_content_type,
    create_watermark_layer,
    empty_relationships_root,
    emu_to_inches,
    emu_to_pt,
    find_font,
    layer_to_png_bytes,
    next_rid,
    qn,
    read_xml_from_zip,
)
from .constants import O_NS, PKG_REL_NS, R_NS, V_NS, W_NS
from .options import WatermarkOptions


def ensure_docx_has_header(input_docx: Path) -> Path:
    try:
        from docx import Document
    except ImportError as e:
        raise ImportError("处理 DOCX 需要安装 python-docx：pip install python-docx") from e

    tmp = Path(tempfile.mkdtemp()) / "with_header.docx"
    doc = Document(str(input_docx))
    for section in doc.sections:
        header = section.header
        if not header.paragraphs:
            header.add_paragraph(" ")
        elif not header.paragraphs[0].text.strip():
            header.paragraphs[0].add_run(" ")
    doc.save(str(tmp))
    return tmp


def get_docx_page_size(input_docx: Path) -> tuple[float, float, float, float]:
    try:
        from docx import Document
    except ImportError as e:
        raise ImportError("处理 DOCX 需要安装 python-docx：pip install python-docx") from e

    doc = Document(str(input_docx))
    section = doc.sections[0]
    width_in = emu_to_inches(section.page_width)
    height_in = emu_to_inches(section.page_height)
    width_pt = emu_to_pt(section.page_width)
    height_pt = emu_to_pt(section.page_height)
    return width_in, height_in, width_pt, height_pt


def remove_old_docx_watermarks(header_root) -> None:
    ns = {"v": V_NS, "w": W_NS}
    for shape in list(header_root.xpath('.//v:shape[starts-with(@id, "UniversalWatermark_")]', namespaces=ns)):
        p = shape
        while p is not None and p.tag != qn(W_NS, "p"):
            p = p.getparent()
        if p is not None and p.getparent() is not None:
            p.getparent().remove(p)


def make_docx_vml_background_paragraph(rel_id: str, page_width_pt: float, page_height_pt: float, shape_id: str):
    from lxml import etree
    nsmap = {"w": W_NS, "r": R_NS, "v": V_NS, "o": O_NS}
    p = etree.Element(qn(W_NS, "p"), nsmap=nsmap)
    r = etree.SubElement(p, qn(W_NS, "r"))
    pict = etree.SubElement(r, qn(W_NS, "pict"))

    shape = etree.SubElement(
        pict,
        qn(V_NS, "shape"),
        {
            "id": shape_id,
            qn(O_NS, "spid"): "_x0000_s2049",
            "type": "#_x0000_t75",
            "style": (
                "position:absolute;"
                "margin-left:0pt;margin-top:0pt;"
                f"width:{page_width_pt:.2f}pt;height:{page_height_pt:.2f}pt;"
                "z-index:-251654144;"
                "mso-position-horizontal:left;"
                "mso-position-horizontal-relative:page;"
                "mso-position-vertical:top;"
                "mso-position-vertical-relative:page;"
                "mso-wrap-edited:f"
            ),
            "filled": "t",
            "stroked": "f",
            qn(O_NS, "allowincell"): "f",
        },
    )

    etree.SubElement(
        shape,
        qn(V_NS, "imagedata"),
        {
            qn(R_NS, "id"): rel_id,
            qn(O_NS, "title"): "universal-watermark",
        },
    )
    return p


def watermark_docx(input_path: Path, output_path: Path, options: WatermarkOptions) -> Path:
    try:
        from lxml import etree
    except ImportError as e:
        raise ImportError("处理 DOCX 需要安装 lxml：pip install lxml") from e

    width_in, height_in, width_pt, height_pt = get_docx_page_size(input_path)
    width_px = int(width_in * options.dpi)
    height_px = int(height_in * options.dpi)
    font = find_font(options.font_path)
    layer = create_watermark_layer(width_px, height_px, options, font)
    png_bytes = layer_to_png_bytes(layer, options.dpi)

    docx_with_header = ensure_docx_has_header(input_path)

    overrides: dict[str, bytes] = {}
    with zipfile.ZipFile(docx_with_header, "r") as zin:
        names = zin.namelist()
        header_files = sorted(
            [n for n in names if re.fullmatch(r"word/header\d+\.xml", n)],
            key=lambda s: int(re.findall(r"\d+", s)[0]),
        )
        if not header_files:
            raise RuntimeError("未找到 DOCX 页眉 XML，无法插入水印。")

        ct_root = read_xml_from_zip(zin, "[Content_Types].xml")
        add_png_content_type(ct_root)
        overrides["[Content_Types].xml"] = etree.tostring(
            ct_root, xml_declaration=True, encoding="UTF-8", standalone=True
        )

        for idx, header_path in enumerate(header_files, start=1):
            media_name = f"word/media/universal_watermark_{idx}.png"
            overrides[media_name] = png_bytes

            rels_path = f"word/_rels/{Path(header_path).name}.rels"
            if rels_path in names:
                rels_root = read_xml_from_zip(zin, rels_path)
            else:
                rels_root = empty_relationships_root()

            rid = next_rid(rels_root)
            etree.SubElement(
                rels_root,
                qn(PKG_REL_NS, "Relationship"),
                {
                    "Id": rid,
                    "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
                    "Target": f"media/universal_watermark_{idx}.png",
                },
            )
            overrides[rels_path] = etree.tostring(
                rels_root, xml_declaration=True, encoding="UTF-8", standalone=True
            )

            header_root = read_xml_from_zip(zin, header_path)
            remove_old_docx_watermarks(header_root)
            wm = make_docx_vml_background_paragraph(rid, width_pt, height_pt, f"UniversalWatermark_{idx}")
            header_root.insert(0, wm)
            overrides[header_path] = etree.tostring(
                header_root, xml_declaration=True, encoding="UTF-8", standalone=True
            )

        with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                name = item.filename
                if name in overrides:
                    zout.writestr(name, overrides.pop(name))
                else:
                    zout.writestr(item, zin.read(name))
            for name, data in overrides.items():
                zout.writestr(name, data)

    return output_path
